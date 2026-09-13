"""Yapılandırılabilir, ATS öncelikli tek sütunlu CV üreticisi.

Önce ``examples/profile.example.json`` dosyasını ``data/profile.json`` olarak kopyalayıp kendi
doğrulanmış bilgilerinle doldur. Bu dosyada örnek kişisel veri bulunmaz.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from job_app.storage import data_dir, data_file, load_json

OUT = data_dir("cv-versions")
PROFILE_FILE = data_file("profile.json")

# Bölüm başlıkları: profil, deneyim, projeler, eğitim, beceriler, diller.
# cv_match bu başlıkları mesleki bölüm olarak tanır; iki modül aynı listeyi
# kullanmazsa üretilen CV'nin bir bölümü yapay zeka yükünden sessizce düşer.
SECTION_LABELS = {
    "TR": ("Profil", "Deneyim", "Projeler", "Eğitim", "Teknik Beceriler", "Diller"),
    "EN": ("Profile", "Experience", "Projects", "Education", "Technical Skills", "Languages"),
}


def load_profile() -> dict:
    if not PROFILE_FILE.exists():
        raise RuntimeError("data/profile.json bulunamadı. examples/profile.example.json dosyasını kopyalayıp kendi gerçek bilgilerinle doldur.")
    data = load_json(PROFILE_FILE, {})
    if not isinstance(data.get("identity"), dict) or not isinstance(data.get("cv_profiles"), dict):
        raise RuntimeError("data/profile.json biçimi geçersiz. examples/profile.example.json şemasını kullan.")
    return data


def font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin, section.bottom_margin = Cm(1.45), Cm(1.35)
    section.left_margin, section.right_margin = Cm(1.55), Cm(1.55)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(2), 1.05
    bullet = doc.styles["List Bullet"]
    bullet.font.name, bullet.font.size = "Arial", Pt(9.6)
    bullet._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    bullet.paragraph_format.left_indent, bullet.paragraph_format.first_line_indent = Cm(0.48), Cm(-0.25)
    bullet.paragraph_format.space_after = Pt(1.5)
    if "CV Section" not in doc.styles:
        section_style = doc.styles.add_style("CV Section", WD_STYLE_TYPE.PARAGRAPH)
        section_style.font.name, section_style.font.size, section_style.font.bold = "Arial", Pt(10.3), True
        section_style.paragraph_format.space_before, section_style.paragraph_format.space_after = Pt(8), Pt(3)


def add_text(paragraph, text: str, size: float = 10, bold: bool = False):
    run = paragraph.add_run(text)
    font(run, size, bold)
    return run


def add_section(doc: Document, title: str) -> None:
    paragraph = doc.add_paragraph(style="CV Section")
    paragraph.paragraph_format.keep_with_next = True
    add_text(paragraph, title.upper(), 10.3, bold=True)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.keep_together = True
    add_text(paragraph, text, 9.6)


def add_entry(doc: Document, entry: dict) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(2), Pt(1)
    paragraph.paragraph_format.keep_with_next = True
    add_text(paragraph, entry["title"], 10.2, bold=True)
    for bullet in entry.get("bullets", []):
        add_bullet(doc, bullet)


def content_for(profile: dict, language: str, focus: str) -> dict:
    variants = profile["cv_profiles"].get(language, {})
    content = variants.get(focus) or variants.get("general")
    if not isinstance(content, dict):
        raise RuntimeError(f"profile.json içinde {language} için '{focus}' veya 'general' CV içeriği yok.")
    return content


def add_header(doc: Document, language: str, identity: dict, content: dict) -> None:
    photo = PROFILE_FILE.parent / identity.get("tr_photo", "")
    use_photo = language == "TR" and photo.is_file()
    header = doc.add_paragraph()
    if use_photo:
        header.paragraph_format.tab_stops.add_tab_stop(Cm(16.4), WD_TAB_ALIGNMENT.RIGHT)
    add_text(header, identity.get("name", "Ad Soyad"), 20, bold=True)
    if use_photo:
        header.add_run("\t").add_picture(str(photo), width=Cm(1.6))
    headline = doc.add_paragraph()
    headline.paragraph_format.space_after = Pt(2)
    add_text(headline, content.get("headline", ""), 11.2, bold=True)
    location = identity.get("tr_address", "") if language == "TR" else identity.get("international_location", "")
    if location:
        add_text(doc.add_paragraph(), location, 8.8)
    contact = " | ".join(value for value in (identity.get("phone", ""), identity.get("email", ""), identity.get("linkedin", ""), identity.get("github", "")) if value)
    if contact:
        add_text(doc.add_paragraph(), contact, 8.8)


def build_cv(language: str, focus: str = "general", filename: str | None = None) -> Path:
    if language not in {"TR", "EN"}:
        raise ValueError("language yalnız TR veya EN olabilir")
    profile = load_profile()
    content = content_for(profile, language, focus)
    doc = Document()
    style_document(doc)
    add_header(doc, language, profile["identity"], content)
    labels = SECTION_LABELS[language]
    add_section(doc, labels[0])
    add_text(doc.add_paragraph(), content.get("summary", ""), 9.8)
    for label, key in ((labels[1], "experience"), (labels[2], "projects")):
        entries = content.get(key, [])
        if entries:
            add_section(doc, label)
            for entry in entries:
                add_entry(doc, entry)
    if content.get("education"):
        add_section(doc, labels[3])
        add_text(doc.add_paragraph(), content["education"], 9.8)
    if content.get("skills"):
        add_section(doc, labels[4])
        for row in content["skills"]:
            paragraph = doc.add_paragraph()
            add_text(paragraph, f"{row['label']}: ", 9.6, bold=True)
            add_text(paragraph, row["value"], 9.6)
    if content.get("language"):
        add_section(doc, labels[5])
        add_text(doc.add_paragraph(), content["language"], 9.6)
    OUT.mkdir(exist_ok=True)
    path = OUT / (filename or f"CV-{language}-TEMEL.docx")
    doc.save(path)
    return path


def main() -> None:
    """Temel TR ve EN CV'lerini üretir."""
    print(build_cv("TR"))
    print(build_cv("EN"))


if __name__ == "__main__":
    main()
