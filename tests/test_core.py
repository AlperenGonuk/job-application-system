"""Çekirdek iş kuralları için ağ/model kullanmayan hızlı kontroller."""
from __future__ import annotations

import importlib.util
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document

import main
from job_app import ai_agents, ai_runner, collect_jobs, cv_document, migration
from job_app import review_detailed as detailed
from job_app import review_initial as initial
from job_app.collect_jobs import matches_preferences, preference_mismatches, preference_sources
from job_app.cv_match import CAREER_SECTION, cv_text
from job_app.dedupe import fingerprint
from job_app.privacy import has_pii, scrub_payload, scrub_text
from job_app.process import hidden_process_options
from job_app.ui_app import cities_for_countries, roles_for_sectors
from job_app.url_guard import URLValidationError, validate_url


class CoreRulesTest(unittest.TestCase):
    def test_candidate_enters_detailed_queue_once_per_mode(self):
        job = {
            "id": "stable-id",
            "company": "Example",
            "title": "Junior Backend Developer",
            "location": "İstanbul",
            "source_url": "https://www.linkedin.com/jobs/view/1",
        }
        key = fingerprint(job)
        original = detailed.all_jobs
        detailed.all_jobs = lambda: [job]
        try:
            review = {key: {"label": "candidate", "needs_detail": True}}
            self.assertEqual(len(detailed.build_queue(review, {}, "flexible")), 1)
            self.assertEqual(len(detailed.build_queue(review, {key: {"mode": "flexible"}}, "flexible")), 0)
            self.assertEqual(len(detailed.build_queue(review, {key: {"mode": "strict"}}, "flexible")), 1)
        finally:
            detailed.all_jobs = original

    def test_out_of_scope_job_never_enters_detailed_queue(self):
        job = {"id": "x", "company": "Example", "title": "Senior Chef", "location": "İzmir",
               "source_url": "https://www.linkedin.com/jobs/view/2"}
        key = fingerprint(job)
        original = detailed.all_jobs
        detailed.all_jobs = lambda: [job]
        try:
            review = {key: {"label": "out_of_scope", "needs_detail": False}}
            self.assertEqual(detailed.build_queue(review, {}, "flexible"), [])
        finally:
            detailed.all_jobs = original

    def test_sensitive_cv_header_is_not_sent_to_ai(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "CV-TR-test.docx"
            document = Document()
            document.add_paragraph("Ada Örnek")
            document.add_paragraph("ada@example.com | +90 555 123 45 67")
            document.add_paragraph("PROJELER")
            document.add_paragraph("Python ile REST API geliştirdim.")
            document.save(path)
            text = cv_text(path)
        self.assertIn("REST API", text)
        self.assertNotIn("Ada Örnek", text)
        self.assertFalse(has_pii(text))

    def test_pii_masking_and_https_guard(self):
        self.assertFalse(has_pii(scrub_text("mail: ada@example.com, +90 555 123 45 67")))
        with self.assertRaises(URLValidationError):
            validate_url("http://localhost:8000/secret")

    def test_role_preference_keeps_relevant_title_variants(self):
        settings = {"target_roles": ["Junior Backend Developer"], "target_cities": ["İstanbul"], "seniority": ["junior"], "remote_ok": True}
        self.assertTrue(matches_preferences({"title": "Junior Python Backend Engineer", "location": "İstanbul"}, settings))
        self.assertFalse(matches_preferences({"title": "Senior Backend Engineer", "location": "İstanbul"}, settings))
        self.assertFalse(matches_preferences({"title": "Junior Backend Engineer", "location": "İzmir"}, settings))

    def test_work_arrangement_filter_respects_remote_choice(self):
        settings = {"target_roles": [], "target_cities": [], "seniority": [], "work_arrangements": ["remote"]}
        self.assertTrue(matches_preferences({"title": "Backend Engineer", "location": "Remote"}, settings))
        self.assertFalse(matches_preferences({"title": "Backend Engineer", "location": "İstanbul"}, settings))

    def test_preference_managed_legacy_source_uses_saved_role_and_location(self):
        config = {"sources": [{"name": "LinkedIn örnek arama", "url": "https://www.linkedin.com/jobs/search?keywords=Junior%20Backend%20Developer"}]}
        settings = {"target_roles": ["Software Engineering Intern", "Data Analyst Intern"], "target_countries": ["Türkiye"], "target_cities": ["Ankara"]}
        generated = preference_sources(config, settings)
        self.assertEqual(len(generated), 2)
        self.assertTrue(all("Ankara%2C%20Turkey" in source["url"] for source in generated))
        self.assertIn("Software%20Engineering%20Intern", generated[0]["url"])

    def test_preference_mismatch_explains_true_conflicts_but_keeps_district(self):
        settings = {"target_roles": ["Software Engineering Intern"], "target_countries": ["Türkiye"], "target_cities": ["Istanbul"], "seniority": [], "work_arrangements": []}
        self.assertEqual(preference_mismatches({"title": "Software Engineer - New Grad", "location": "Sarıyer, Turkey"}, settings), [])
        reasons = preference_mismatches({"title": "Senior Accountant", "location": "Ankara, Turkey"}, settings)
        self.assertIn("role", reasons)
        self.assertIn("city", reasons)

    def test_country_and_city_filters_support_global_options_and_turkey_alias(self):
        turkey = {"target_roles": [], "target_countries": ["Türkiye"], "target_cities": ["All cities"], "seniority": [], "work_arrangements": []}
        self.assertTrue(matches_preferences({"title": "Intern", "location": "Istanbul, Turkey"}, turkey))
        world = {"target_roles": [], "target_countries": ["Worldwide"], "target_cities": ["All cities"], "seniority": [], "work_arrangements": []}
        self.assertTrue(matches_preferences({"title": "Intern", "location": "Tokyo, Japan"}, world))

    def test_sector_and_country_confirmation_sources_are_dependent(self):
        technology_roles = roles_for_sectors({"Technology / SaaS"})
        self.assertIn("Software Engineering Intern", technology_roles)
        self.assertNotIn("Legal Intern", technology_roles)
        turkey_cities = cities_for_countries({"Türkiye"})
        self.assertIn("All cities in Türkiye", turkey_cities)
        self.assertIn("İstanbul" if "İstanbul" in turkey_cities else "Istanbul", turkey_cities)
        self.assertNotIn("Tokyo", turkey_cities)

    def test_public_card_parser_uses_only_card_fields(self):
        page = '''<div class="base-search-card"><h3 class="base-search-card__title">Junior Backend Engineer</h3><h4 class="base-search-card__subtitle"><a>Example Ltd</a></h4><span class="job-search-card__location">Istanbul, Turkey</span><a class="base-card__full-link" href="https://www.linkedin.com/jobs/view/42"></a><time datetime="2026-09-09"></time></li>'''
        with patch.object(collect_jobs, "safe_fetch", return_value=page):
            cards = collect_jobs.linkedin_cards({"name": "Example", "url": "https://www.linkedin.com/jobs/search"})
        self.assertEqual(len(cards), 1)
        self.assertEqual(cards[0]["company"], "Example Ltd")
        self.assertEqual(cards[0]["source_url"], "https://www.linkedin.com/jobs/view/42")

    def test_first_run_creates_private_source_config_from_template(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            target, template = root / "data" / "job-sources.json", root / "job-sources.example.json"
            template.write_text(json.dumps({"max_raw_cards": 5, "sources": []}), encoding="utf-8")
            with patch.object(collect_jobs, "SOURCES_FILE", target), patch.object(collect_jobs, "SOURCES_TEMPLATE", template):
                self.assertEqual(collect_jobs.load_source_config()["max_raw_cards"], 5)
            self.assertTrue(target.exists())

    def test_initial_review_persists_only_valid_mocked_result(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            history, results = root / "history", root / "results"
            history.mkdir()
            results.mkdir()
            job = {"id": "stable-job", "company": "Example", "title": "Junior Backend Engineer", "location": "İstanbul", "source_url": "https://www.linkedin.com/jobs/view/42", "published_at": "2026-09-09"}
            (history / "scan-1.json").write_text(json.dumps({"new_jobs": [job]}), encoding="utf-8")
            state_file = root / "state.json"
            model_output = json.dumps([{"id": "stable-job", "label": "candidate", "needs_detail": True}])
            with patch.object(initial, "HISTORY_DIR", history), patch.object(initial, "RESULT_DIR", results), patch.object(initial, "STATE_FILE", state_file), patch.object(initial, "is_agent_available", return_value=True), patch.object(initial, "run_agent", return_value=model_output), patch.object(sys, "argv", ["initial-review", "--limit", "1"]):
                with redirect_stdout(io.StringIO()):
                    initial.main()
            persisted = json.loads(state_file.read_text(encoding="utf-8"))
            self.assertEqual(persisted[fingerprint(job)]["label"], "candidate")
            self.assertIs(persisted[fingerprint(job)]["needs_detail"], True)

    def test_initial_review_rejects_out_of_schema_output(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            history, results = root / "history", root / "results"
            history.mkdir()
            results.mkdir()
            job = {"id": "stable-job", "company": "Example", "title": "Junior Backend Engineer", "location": "İstanbul", "source_url": "https://www.linkedin.com/jobs/view/42", "published_at": "2026-09-09"}
            (history / "scan-1.json").write_text(json.dumps({"new_jobs": [job]}), encoding="utf-8")
            state_file = root / "state.json"
            model_output = json.dumps([{"id": "stable-job", "label": "aday", "needs_detail": "evet"}])
            with patch.object(initial, "HISTORY_DIR", history), patch.object(initial, "RESULT_DIR", results), patch.object(initial, "STATE_FILE", state_file), patch.object(initial, "is_agent_available", return_value=True), patch.object(initial, "run_agent", return_value=model_output), patch.object(sys, "argv", ["initial-review", "--limit", "1"]):
                with self.assertRaises(ValueError):
                    initial.main()
            self.assertFalse(state_file.exists())

    def test_candidate_facts_excludes_identity_even_if_user_added_it(self):
        with tempfile.TemporaryDirectory() as directory:
            evidence = Path(directory) / "evidence.json"
            evidence.write_text(json.dumps({"name": "Ada Örnek", "email": "ada@example.com", "projects": ["Python API"]}), encoding="utf-8")
            with patch.object(detailed, "EVIDENCE_FILE", evidence):
                self.assertEqual(detailed.candidate_facts(), {"projects": ["Python API"]})

    def test_detailed_review_stops_before_fetching_when_evidence_missing(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with patch.object(detailed, "EVIDENCE_FILE", root / "missing.json"), \
                 patch.object(detailed, "LOCK_FILE", root / ".lock"), \
                 patch.object(detailed, "MATCH_FILE", root / "state.json"), \
                 patch.object(detailed, "fetch_description") as fetch, \
                 patch.object(sys, "argv", ["review_detailed"]):
                with self.assertRaises(RuntimeError):
                    detailed.main()
                fetch.assert_not_called()
                self.assertFalse((root / ".lock").exists())
                self.assertFalse((root / "state.json").exists())


class AgentRegistryTest(unittest.TestCase):
    def test_argument_agent_puts_prompt_on_command_line(self):
        command, stdin_text = ai_agents.build_command("hermes", "merhaba", task="fast", settings={})
        self.assertIsNone(stdin_text)
        self.assertIn("merhaba", command)
        self.assertIn("claude-haiku-4-5-20251001", command)

    def test_stdin_agent_keeps_prompt_out_of_command_line(self):
        command, stdin_text = ai_agents.build_command("claude", "merhaba", task="deep", settings={})
        self.assertEqual(stdin_text, "merhaba")
        self.assertNotIn("merhaba", command)
        self.assertIn("sonnet", command)

    def test_empty_model_drops_the_model_flag(self):
        command, _ = ai_agents.build_command("agy", "merhaba", task="fast", settings={})
        self.assertNotIn("--model", command)

    def test_user_model_override_wins_over_profile_default(self):
        command, _ = ai_agents.build_command("hermes", "merhaba", task="fast", settings={"ai_model_fast": "kendi-modelim"})
        self.assertIn("kendi-modelim", command)
        self.assertNotIn("claude-haiku-4-5-20251001", command)

    def test_recommended_models_cover_fast_and_deep_for_known_agents(self):
        for agent, models in ai_agents.RECOMMENDED_MODELS.items():
            self.assertIn(agent, ai_agents.AGENT_PROFILES, agent)
            self.assertEqual(set(models), {"fast", "deep"}, agent)
            self.assertTrue(all(name.strip() for name in models.values()), agent)

    def test_unknown_agent_has_no_invented_recommendation(self):
        # Doğrulanmamış model adı uydurulmamalı; boş sözlük arayüzde genel öneriye düşer.
        self.assertEqual(ai_agents.recommended_models("gemini"), {})
        self.assertEqual(ai_agents.recommended_models("cursor-agent"), {})

    def test_list_models_commands_point_at_known_agents(self):
        for agent in ai_agents.LIST_MODELS_COMMAND:
            self.assertIn(agent, ai_agents.AGENT_PROFILES, agent)

    def test_custom_command_substitutes_prompt_placeholder(self):
        settings = {"ai_agent": "custom", "ai_custom_command": "myagent run --text {prompt}"}
        command, stdin_text = ai_agents.build_command("custom", "merhaba", task="fast", settings=settings)
        self.assertEqual(command, ["myagent", "run", "--text", "merhaba"])
        self.assertIsNone(stdin_text)

    def test_custom_command_without_placeholder_uses_stdin(self):
        settings = {"ai_agent": "custom", "ai_custom_command": "myagent run"}
        command, stdin_text = ai_agents.build_command("custom", "merhaba", task="fast", settings=settings)
        self.assertEqual(command, ["myagent", "run"])
        self.assertEqual(stdin_text, "merhaba")

    def test_explicit_but_missing_agent_does_not_silently_fall_back(self):
        with patch.object(ai_agents, "is_installed", return_value=False):
            self.assertIsNone(ai_agents.resolve_agent({"ai_agent": "codex"}))

    def test_auto_choice_picks_first_installed_agent(self):
        with patch.object(ai_agents, "available_agents", return_value=["agy", "pi"]):
            self.assertEqual(ai_agents.resolve_agent({"ai_agent": "auto"}), "agy")


class PrivacyTest(unittest.TestCase):
    def test_public_job_url_survives_scrubbing(self):
        # LinkedIn ilan numarası 10 hanelidir; telefon sanılıp maskelenirse
        # modelin döndürdüğü id hiçbir ilanla eşleşmez.
        url = "https://tr.linkedin.com/jobs/view/game-developer-at-loop-games-4452573928"
        cleaned = scrub_text(f"{url} iletisim: ada@example.com +90 555 123 45 67")
        self.assertIn(url, cleaned)
        self.assertIn("TELEFON GİZLİ", cleaned)
        self.assertIn("E-POSTA GİZLİ", cleaned)
        self.assertFalse(has_pii(cleaned))

    def test_address_in_experience_bullets_is_masked(self):
        payload = {"experience": [{"title": "Backend Stajyeri", "bullets": [
            "Adres: Test Mah. Deneme Sok. No: 4",
            "Ofis ziyareti: Test Mah. Deneme Sok. No: 4 Kadıköy",
            "Python ile REST API geliştirdim ve PostgreSQL sorgularını hızlandırdım.",
        ]}]}
        cleaned = scrub_payload(payload)
        labeled, unlabeled, professional = cleaned["experience"][0]["bullets"]
        for text in (labeled, unlabeled):
            self.assertNotIn("Test Mah", text)
            self.assertNotIn("Deneme Sok", text)
            self.assertNotIn("No: 4", text)
            self.assertIn("[ADRES GİZLİ]", text)
            self.assertFalse(has_pii(text))
        self.assertEqual(professional, payload["experience"][0]["bullets"][2])
        self.assertEqual(cleaned["experience"][0]["title"], "Backend Stajyeri")
        self.assertTrue(has_pii("Adres: Test Mah. Deneme Sok. No: 4"))

    def test_personal_url_parameters_are_masked(self):
        url = "https://example.com/?email=test@example.com&phone=05550000000"
        self.assertTrue(has_pii(url))
        cleaned = scrub_text(f"Başvuru bağlantısı: {url}")
        self.assertNotIn("test@example.com", cleaned)
        self.assertNotIn("05550000000", cleaned)
        self.assertIn("https://example.com/?email=", cleaned)
        self.assertFalse(has_pii(cleaned))
        self.assertNotIn("ada@example.com", scrub_text("https://example.com/u/ada%40example.com#ref"))
        self.assertNotIn("ada", scrub_text("https://ada:secret@example.com/jobs/1"))

    def test_job_url_identifiers_are_preserved_exactly(self):
        for url in (
            "https://www.linkedin.com/jobs/view/4452573928/?refId=abc%3D%3D&trackingId=x%2By&currentJobId=4452573928",
            "https://boards.greenhouse.io/acme/jobs/5551234567?gh_jid=5551234567",
            "https://jobs.lever.co/acme/1b2c3d4e-0000-4000-8000-123456789abc?lever-source=LinkedIn",
            "https://www.kariyer.net/is-ilani/acme-backend-developer-3942211#apply",
        ):
            self.assertEqual(scrub_text(url), url)
            self.assertFalse(has_pii(url))


NPM_SHIM = """@ECHO off\r
GOTO start\r
:find_dp0\r
SET dp0=%~dp0\r
EXIT /b\r
:start\r
SETLOCAL\r
CALL :find_dp0\r
\r
IF EXIST "%dp0%\\node.exe" (\r
  SET "_prog=%dp0%\\node.exe"\r
) ELSE (\r
  SET "_prog=node"\r
  SET PATHEXT=%PATHEXT:;.JS;=;%\r
)\r
\r
endLocal & goto #_undefined_# 2>NUL || title %COMSPEC% & "%_prog%"  "%dp0%\\node_modules\\fake-agent\\cli.js" %*\r
"""
ECHO_ARGV_SCRIPT = "process.stdout.write(JSON.stringify(process.argv.slice(2)));\n"


def make_npm_shim(directory: Path, name: str) -> Path:
    """Gerçek npm cmd-shim biçiminde, argümanları JSON olarak basan sahte ajan."""
    script = directory / "node_modules" / "fake-agent" / "cli.js"
    script.parent.mkdir(parents=True, exist_ok=True)
    script.write_text(ECHO_ARGV_SCRIPT, encoding="utf-8")
    shim = directory / f"{name}.cmd"
    shim.write_bytes(NPM_SHIM.encode("ascii"))
    return shim


class AgentTransportTest(unittest.TestCase):
    """Windows .cmd kısayollarında çok satırlı ve özel karakterli istem taşıma."""

    HOSTILE_JOB_TEXT = (
        "İlan başlığı: Backend Developer\n"
        "JOB_TEXT_SENTINEL \"tırnak\" & echo PWNED> pwned.txt & | < > ^ %PATH% !x!\n"
        "Son satır"
    )

    @unittest.skipUnless(os.name == "nt", "Windows .cmd davranışı")
    def test_old_style_shim_to_native_exe_is_resolved(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            target = root / "node_modules" / "native" / "bin" / "tool.exe"
            target.parent.mkdir(parents=True)
            target.write_bytes(b"")
            shim = root / "tool.cmd"
            shim.write_text('@"%~dp0\\node_modules\\native\\bin\\tool.exe"   %*\r\n', encoding="ascii")
            self.assertEqual(ai_agents.resolve_batch_shim(shim), [str(target)])

    def test_unrecognized_batch_file_is_not_resolved(self):
        with tempfile.TemporaryDirectory() as directory:
            shim = Path(directory) / "agent.cmd"
            shim.write_text("@echo off\r\necho ARGS: %*\r\n", encoding="ascii")
            self.assertIsNone(ai_agents.resolve_batch_shim(shim))

    @unittest.skipIf(os.name == "nt", "POSIX kabuk tırnak kuralları")
    def test_custom_command_accepts_single_quoted_path_on_posix(self):
        with tempfile.TemporaryDirectory(prefix="agent dir ") as directory:
            program = Path(directory) / "my agent"
            settings = {"ai_agent": "custom", "ai_custom_command": f"'{program}' run {{prompt}}"}
            command, stdin_text = ai_agents.build_command("custom", "iki kelime", task="fast", settings=settings)
            self.assertEqual(command, [str(program), "run", "iki kelime"])
            self.assertIsNone(stdin_text)

    @unittest.skipUnless(os.name == "nt", "Windows komut satırı ayrıştırması")
    def test_custom_command_keeps_backslashes_in_quoted_windows_path(self):
        settings = {"ai_agent": "custom", "ai_custom_command": r'"C:\Agent Tools\agent.exe" exec {prompt}'}
        command, _ = ai_agents.build_command("custom", "iki kelime", task="fast", settings=settings)
        self.assertEqual(command, [r"C:\Agent Tools\agent.exe", "exec", "iki kelime"])

    @unittest.skipUnless(os.name == "nt", "Windows .cmd davranışı")
    @unittest.skipUnless(shutil.which("node"), "node gerekli")
    def test_run_agent_delivers_full_multiline_prompt_through_npm_shim(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            make_npm_shim(root, "gemini")
            previous = os.getcwd()
            os.chdir(root)
            try:
                with patch.dict(os.environ, {"PATH": f"{root}{os.pathsep}{os.environ.get('PATH', '')}"}), \
                        patch.object(ai_runner, "LOCK_FILE", root / ".lock"):
                    output = ai_runner.run_agent(self.HOSTILE_JOB_TEXT, settings={"ai_agent": "gemini"})
            finally:
                os.chdir(previous)
            arguments = json.loads(output)
            self.assertEqual(arguments[0], "-p")
            self.assertEqual(arguments[1], ai_runner.NO_TOOL_PREAMBLE + scrub_text(self.HOSTILE_JOB_TEXT))
            self.assertIn("JOB_TEXT_SENTINEL", arguments[1])
            self.assertTrue(arguments[1].endswith("Son satır"))
            self.assertFalse((root / "pwned.txt").exists())

    def test_safe_argument_whitelist_rejects_trailing_newline(self):
        self.assertIsNotNone(ai_agents._CMD_SAFE_ARGUMENT.fullmatch("claude-sonnet-4-6"))
        for argument in ("sonnet\n", "sonnet\r\n", "a\nb", "x & y"):
            self.assertIsNone(ai_agents._CMD_SAFE_ARGUMENT.fullmatch(argument), repr(argument))

    @unittest.skipUnless(os.name == "nt", "Windows .cmd davranışı")
    def test_batch_with_trailing_newline_argument_is_refused(self):
        with tempfile.TemporaryDirectory() as directory:
            shim = Path(directory) / "agent.cmd"
            shim.write_text("@echo off\r\necho ARGS: %*\r\n", encoding="ascii")
            with self.assertRaises(ai_agents.CommandTransportError):
                ai_agents.safe_command([str(shim), "-m", "sonnet\n"])

    @unittest.skipUnless(os.name == "nt", "Windows .cmd davranışı")
    def test_run_agent_refuses_plain_batch_instead_of_truncating(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "gemini.cmd").write_text('@echo off\r\necho ran> "%~dp0ran.txt"\r\necho ARGS: %*\r\n', encoding="ascii")
            with patch.dict(os.environ, {"PATH": f"{root}{os.pathsep}{os.environ.get('PATH', '')}"}), \
                    patch.object(ai_runner, "LOCK_FILE", root / ".lock"):
                with self.assertRaises(ai_runner.AgentError) as raised:
                    ai_runner.run_agent(self.HOSTILE_JOB_TEXT, settings={"ai_agent": "gemini"})
            self.assertIn("gemini.cmd", str(raised.exception).lower())
            self.assertFalse((root / "ran.txt").exists())

    @unittest.skipUnless(os.name == "nt", "Windows .cmd davranışı")
    @unittest.skipUnless(shutil.which("node"), "node gerekli")
    def test_custom_batch_command_with_placeholder_bypasses_cmd(self):
        with tempfile.TemporaryDirectory() as directory:
            tools = Path(directory) / "my tools"
            shim = make_npm_shim(tools, "myagent")
            settings = {"ai_agent": "custom", "ai_custom_command": f'"{shim}" --text {{prompt}}'}
            command, stdin_text = ai_agents.build_command("custom", self.HOSTILE_JOB_TEXT, task="fast", settings=settings)
            self.assertIsNone(stdin_text)
            self.assertTrue(command[1].endswith("cli.js"))
            self.assertEqual(command[2:], ["--text", self.HOSTILE_JOB_TEXT])
            result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", cwd=directory)
            self.assertEqual(json.loads(result.stdout), ["--text", self.HOSTILE_JOB_TEXT])
            self.assertFalse((Path(directory) / "pwned.txt").exists())

    @unittest.skipUnless(os.name == "nt", "Windows .cmd davranışı")
    def test_custom_batch_command_without_placeholder_keeps_stdin(self):
        with tempfile.TemporaryDirectory() as directory:
            shim = Path(directory) / "plain.cmd"
            shim.write_text("@echo off\r\nmore\r\n", encoding="ascii")
            settings = {"ai_agent": "custom", "ai_custom_command": f"{shim} run"}
            command, stdin_text = ai_agents.build_command("custom", self.HOSTILE_JOB_TEXT, task="fast", settings=settings)
            self.assertEqual(command, [str(shim), "run"])
            self.assertEqual(stdin_text, self.HOSTILE_JOB_TEXT)


class BuildExeTest(unittest.TestCase):
    def test_every_dynamic_task_module_is_bundled(self):
        spec = importlib.util.spec_from_file_location("build_exe_under_test", ROOT / "scripts" / "build_exe.py")
        build_exe = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(build_exe)
        command = build_exe.pyinstaller_command()
        hidden = {command[index + 1] for index, part in enumerate(command) if part == "--hidden-import"}
        self.assertLessEqual(set(main.TASKS.values()), hidden)
        for module in ("job_app.collect_jobs", "job_app.cv_for_job", "job_app.cv_match", "job_app.cv_document"):
            self.assertIn(module, hidden)
        for module in main.TASKS.values():
            self.assertIsNotNone(importlib.util.find_spec(module), module)
        self.assertEqual(command[command.index("--collect-submodules") + 1], "job_app")
        self.assertEqual(command[-1], str(ROOT / "main.py"))


class CvSectionTest(unittest.TestCase):
    def test_generator_headings_are_recognized_as_career_sections(self):
        for labels in cv_document.SECTION_LABELS.values():
            for label in labels:
                self.assertTrue(CAREER_SECTION.match(label.upper()), label)

    def test_generated_cv_summary_is_extracted_in_both_languages(self):
        def content(summary: str) -> dict:
            return {"headline": "Backend Developer", "summary": summary,
                    "experience": [{"title": "Intern", "bullets": ["Built REST APIs with Python."]}],
                    "education": "Computer Engineering", "skills": [{"label": "Backend", "value": "Python, FastAPI"}],
                    "language": "English (C1)"}

        profile = {
            "identity": {"name": "Ada Örnek", "email": "ada@example.com", "phone": "+90 555 123 45 67",
                         "tr_address": "Test Mah. Deneme Sok. No: 4", "international_location": "Istanbul, Türkiye"},
            "cv_profiles": {"EN": {"general": content("EN_SUMMARY_SENTINEL builds reliable services.")},
                            "TR": {"general": content("TR_SUMMARY_SENTINEL güvenilir servisler geliştirir.")}},
        }
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            profile_file = root / "profile.json"
            profile_file.write_text(json.dumps(profile, ensure_ascii=False), encoding="utf-8")
            with patch.object(cv_document, "OUT", root / "cv-versions"), patch.object(cv_document, "PROFILE_FILE", profile_file):
                texts = {language: cv_text(cv_document.build_cv(language)) for language in ("EN", "TR")}
        self.assertIn("PROFILE", texts["EN"])
        self.assertIn("EN_SUMMARY_SENTINEL", texts["EN"])
        self.assertIn("TR_SUMMARY_SENTINEL", texts["TR"])
        for text in texts.values():
            self.assertIn("Built REST APIs", text)
            self.assertIn("Python, FastAPI", text)
            self.assertNotIn("Ada Örnek", text)
            self.assertNotIn("Deneme", text)
            self.assertFalse(has_pii(text))


class HiddenProcessTest(unittest.TestCase):
    def test_windows_gets_no_console_window(self):
        options = hidden_process_options()
        if sys.platform.startswith("win"):
            self.assertEqual(options["creationflags"], subprocess.CREATE_NO_WINDOW)
            self.assertIn("startupinfo", options)
        else:
            self.assertEqual(options, {})


class MigrationTest(unittest.TestCase):
    def test_turkish_keys_and_values_become_english(self):
        old = {
            "url:x": {"etiket": "aday", "sonnet": "evet"},
            "url:y": {"karar": "başvurma", "uyum_puani": 40, "cv_tipi": "genel", "cv_dili": "TR",
                      "gerekce": "kısa not", "mode": "cok_esnek"},
        }
        new = migration.translate(old)
        self.assertEqual(new["url:x"], {"label": "candidate", "needs_detail": True})
        self.assertEqual(new["url:y"]["decision"], "skip")
        self.assertEqual(new["url:y"]["match_score"], 40)
        self.assertEqual(new["url:y"]["cv_focus"], "general")
        self.assertEqual(new["url:y"]["cv_language"], "TR")
        self.assertEqual(new["url:y"]["reason"], "kısa not")
        self.assertEqual(new["url:y"]["mode"], "very_flexible")

    def test_free_text_is_not_translated_by_accident(self):
        # "başvur" serbest metin alanında geçse bile yalnız karar alanı çevrilir.
        translated = migration.translate({"gerekce": "başvur", "karar": "başvur"})
        self.assertEqual(translated["reason"], "başvur")
        self.assertEqual(translated["decision"], "apply")

    def test_work_arrangements_list_is_translated(self):
        translated = migration.translate({"work_arrangements": ["ofis", "uzaktan"], "detailed_review_mode": "esnek"})
        self.assertEqual(translated["work_arrangements"], ["onsite", "remote"])
        self.assertEqual(translated["detailed_review_mode"], "flexible")

    def test_profile_cv_variant_is_renamed(self):
        profile = {"cv_profiles": {"TR": {"genel": {"headline": "x"}, "python_backend": {"headline": "y"}}}}
        renamed = migration.translate_profile(profile)
        self.assertIn("general", renamed["cv_profiles"]["TR"])
        self.assertIn("python_backend", renamed["cv_profiles"]["TR"])
        self.assertNotIn("genel", renamed["cv_profiles"]["TR"])

    def test_migration_end_to_end_updates_generated_cv_paths(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            data = root / "data"
            legacy = data / "CV-Sürümleri"
            root_legacy = root / "CV-Sürümleri"
            versions = data / "cv-versions"
            for folder in (legacy, root_legacy, versions):  # uygulama boş hedefi önceden açabilir
                folder.mkdir(parents=True)
            (legacy / "CV-EN-Acme-Backend.docx").write_bytes(b"synthetic-a")
            (root_legacy / "CV-TR-Beta-Data.docx").write_bytes(b"synthetic-b")
            (legacy / "CV-TR-Same.docx").write_bytes(b"old-conflict")
            (versions / "CV-TR-Same.docx").write_bytes(b"new-conflict")
            outside = root / "elsewhere" / "CV-TR-Other.docx"
            outside.parent.mkdir()
            outside.write_bytes(b"synthetic-c")
            selections = {
                "url:a": {"cv": "CV-EN-Acme-Backend", "generated_file": str(legacy / "CV-EN-Acme-Backend.docx"),
                          "match_summary": {"karar": "başvur"}},
                "url:b": {"cv": "CV-TR-Beta-Data", "generated_file": str(root_legacy / "CV-TR-Beta-Data.docx")},
                "url:c": {"generated_file": str(outside)},
                "url:d": {"generated_file": str(legacy / ".." / ".." / "elsewhere" / "CV-TR-Other.docx")},
                "url:e": {"cv": "CV-TR-TEMEL", "language": "TR"},
                "url:f": {"generated_file": str(legacy / "CV-TR-Same.docx")},
            }
            (data / "arayuz-cv-secimleri.json").write_text(json.dumps(selections, ensure_ascii=False), encoding="utf-8")

            with patch.object(migration, "ROOT", root), patch.object(migration, "DATA_DIR", data), \
                    patch.object(migration, "MARKER_FILE", data / ".migration-done"):
                summary = migration.run()

            migrated = json.loads((data / "cv-selections.json").read_text(encoding="utf-8"))
            self.assertEqual(summary["path_references"], 2)
            for key, name in (("url:a", "CV-EN-Acme-Backend.docx"), ("url:b", "CV-TR-Beta-Data.docx")):
                path = Path(migrated[key]["generated_file"])
                self.assertEqual(path, versions / name)
                self.assertTrue(path.is_file())
                # Arayüzün "CV'yi aç" güvenlik kontrolüyle aynı koşul
                path.resolve().relative_to(versions.resolve())
            self.assertEqual(migrated["url:a"]["match_summary"], {"decision": "apply"})
            self.assertEqual(migrated["url:c"]["generated_file"], selections["url:c"]["generated_file"])
            self.assertEqual(migrated["url:d"]["generated_file"], selections["url:d"]["generated_file"])
            self.assertEqual(migrated["url:e"], selections["url:e"])
            # Çakışmada eski dosya yerinde kalır; kayıt başka bir dosyaya çevrilmez.
            self.assertEqual(migrated["url:f"]["generated_file"], selections["url:f"]["generated_file"])
            self.assertEqual((versions / "CV-TR-Same.docx").read_bytes(), b"new-conflict")
            self.assertFalse(root_legacy.exists())

    def test_already_migrated_install_repairs_generated_cv_paths_idempotently(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            data = root / "data"
            legacy = data / "CV-Sürümleri"
            root_legacy = root / "CV-Sürümleri"
            versions = data / "cv-versions"
            versions.mkdir(parents=True)
            legacy.mkdir()  # çakışan eski dosya yerinde kalmış
            # Önceki taşıma dosyaları taşımış ama yolları güncellememiş.
            (versions / "CV-EN-Acme-Backend.docx").write_bytes(b"moved-a")
            (versions / "CV-TR-Beta-Data.docx").write_bytes(b"moved-b")
            (legacy / "CV-TR-Same.docx").write_bytes(b"old-conflict")
            (versions / "CV-TR-Same.docx").write_bytes(b"new-conflict")
            outside = root / "elsewhere" / "CV-TR-Other.docx"
            outside.parent.mkdir()
            outside.write_bytes(b"outside")
            (data / ".migration-done").write_text("{}", encoding="utf-8")
            # Tam taşımanın yeniden çalışmadığını kanıtlayan eski veriler:
            (data / "profil.json").write_text(json.dumps({"cv_profiles": {}}), encoding="utf-8")
            selections = {
                "url:a": {"generated_file": str(legacy / "CV-EN-Acme-Backend.docx"), "match_summary": {"karar": "başvur"}},
                "url:b": {"generated_file": str(root_legacy / "CV-TR-Beta-Data.docx")},
                "url:c": {"generated_file": str(outside)},
                "url:d": {"generated_file": str(legacy / ".." / ".." / "elsewhere" / "CV-TR-Other.docx")},
                "url:f": {"generated_file": str(legacy / "CV-TR-Same.docx")},
                "url:g": {"generated_file": str(legacy / "CV-EN-Missing.docx")},
            }
            selections_file = data / "cv-selections.json"
            selections_file.write_text(json.dumps(selections, ensure_ascii=False), encoding="utf-8")

            with patch.object(migration, "ROOT", root), patch.object(migration, "DATA_DIR", data), \
                    patch.object(migration, "MARKER_FILE", data / ".migration-done"):
                first = migration.run()
                repaired_bytes = selections_file.read_bytes()
                repaired_mtime = selections_file.stat().st_mtime_ns
                second = migration.run()

            self.assertEqual(first, {"skipped": True, "path_references": 2})
            self.assertEqual(second, {"skipped": True, "path_references": 0})
            self.assertEqual(selections_file.read_bytes(), repaired_bytes)
            self.assertEqual(selections_file.stat().st_mtime_ns, repaired_mtime)
            repaired = json.loads(repaired_bytes.decode("utf-8"))
            self.assertEqual(Path(repaired["url:a"]["generated_file"]), versions / "CV-EN-Acme-Backend.docx")
            self.assertEqual(Path(repaired["url:b"]["generated_file"]), versions / "CV-TR-Beta-Data.docx")
            for key in ("url:c", "url:d", "url:f", "url:g"):
                self.assertEqual(repaired[key]["generated_file"], selections[key]["generated_file"], key)
            # Diğer kullanıcı verisine dokunulmadı: çeviri ve dosya adı taşıma yok.
            self.assertEqual(repaired["url:a"]["match_summary"], {"karar": "başvur"})
            self.assertTrue((data / "profil.json").exists())
            self.assertFalse((data / "profile.json").exists())
            self.assertEqual((legacy / "CV-TR-Same.docx").read_bytes(), b"old-conflict")


if __name__ == "__main__":
    unittest.main()
